"""
job_pipeline.py — AI-Powered Job Application Pipeline
Searches LinkedIn & Naukri, runs Claude fitment analysis,
exports a ranked Excel report, and opens top matches in browser.
"""

import re
import time
import webbrowser
import anthropic
from datetime import datetime
from playwright.sync_api import sync_playwright
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from rich.console import Console
from rich.progress import Progress, SpinnerColumn, TextColumn
from rich.table import Table

from config import (
    ANTHROPIC_KEY, RESUME, JOB_TITLE, LOCATIONS,
    MAX_JOBS, TOP_N_OPEN, MIN_FIT_PCT,
    SEARCH_LINKEDIN, SEARCH_NAUKRI
)

console = Console()

# ─── SCRAPERS ────────────────────────────────────────────────────────────────

def scrape_linkedin(page, title, location, max_results=8):
    jobs = []
    url = (
        f"https://www.linkedin.com/jobs/search/"
        f"?keywords={title.replace(' ', '%20')}"
        f"&location={location.replace(' ', '%20')}"
        f"&f_TPR=r604800"
    )
    console.log(f"[cyan]LinkedIn → {location}:[/cyan] {url}")
    try:
        page.goto(url, timeout=20000)
        page.wait_for_timeout(3000)
        for card in page.query_selector_all(".job-search-card")[:max_results]:
            try:
                jobs.append({
                    "portal":   "LinkedIn",
                    "title":    _text(card, ".base-search-card__title") or title,
                    "company":  _text(card, ".base-search-card__subtitle") or "Unknown",
                    "location": _text(card, ".job-search-card__location") or location,
                    "url":      _href(card, "a.base-card__full-link") or url,
                    "jd":       ""
                })
            except Exception:
                continue
    except Exception as e:
        console.log(f"[yellow]LinkedIn warning ({location}):[/yellow] {e}")
    return jobs


def scrape_naukri(page, title, location, max_results=8):
    jobs = []
    url = (
        f"https://www.naukri.com/"
        f"{title.lower().replace(' ', '-')}-jobs-in-"
        f"{location.lower().replace(' ', '-')}"
    )
    console.log(f"[cyan]Naukri → {location}:[/cyan] {url}")
    try:
        page.goto(url, timeout=20000)
        page.wait_for_timeout(3000)
        for card in page.query_selector_all(".srp-jobtuple-wrapper, article.jobTuple")[:max_results]:
            try:
                jobs.append({
                    "portal":   "Naukri",
                    "title":    _text(card, ".title, .jobTitle") or title,
                    "company":  _text(card, ".comp-name, .companyName") or "Unknown",
                    "location": _text(card, ".locWdth, .location") or location,
                    "url":      _href(card, "a.title, a.jobTitle") or url,
                    "jd":       ""
                })
            except Exception:
                continue
    except Exception as e:
        console.log(f"[yellow]Naukri warning ({location}):[/yellow] {e}")
    return jobs


def fetch_jd(page, url):
    try:
        page.goto(url, timeout=15000)
        page.wait_for_timeout(2000)
        for selector in [
            ".show-more-less-html__markup",
            ".description__text",
            ".job-desc",
            ".JDC",
            "#job-desc",
            "main",
            "article"
        ]:
            el = page.query_selector(selector)
            if el:
                return el.inner_text().strip()[:3000]
    except Exception:
        pass
    return ""


def _text(el, selector):
    found = el.query_selector(selector)
    return found.inner_text().strip() if found else None

def _href(el, selector):
    found = el.query_selector(selector)
    return found.get_attribute("href") if found else None

# ─── CLAUDE FITMENT ──────────────────────────────────────────────────────────

def run_fitment(client, job):
    jd = job["jd"] or (
        f"Senior Project Manager at {job['company']} in {job['location']}. "
        "Requires project management, stakeholder management, team leadership."
    )
    prompt = f"""You are a resume fitment analyzer. Reply ONLY in this exact format, nothing else:

FIT:{{number 0-100}}
MATCHES:skill1,skill2,skill3
GAPS:gap1,gap2,gap3
VERDICT:one concise sentence

Resume:
{RESUME.strip()}

Job Title: {job['title']}
Company:   {job['company']}
JD:
{jd[:2000]}"""

    try:
        msg = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=300,
            messages=[{"role": "user", "content": prompt}]
        )
        r = msg.content[0].text
        return {
            "fit":     int((re.search(r"FIT:(\d+)", r) or [None, 0])[1]),
            "matches": (re.search(r"MATCHES:(.+)", r) or [None, ""])[1].strip(),
            "gaps":    (re.search(r"GAPS:(.+)",    r) or [None, ""])[1].strip(),
            "verdict": (re.search(r"VERDICT:(.+)", r) or [None, ""])[1].strip(),
        }
    except Exception as e:
        return {"fit": 0, "matches": "", "gaps": "", "verdict": f"Error: {e}"}

# ─── EXCEL REPORT ────────────────────────────────────────────────────────────

def build_excel(results, path):
    wb  = Workbook()
    ws  = wb.active
    ws.title = "Fitment Report"

    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap   = Alignment(vertical="top", wrap_text=True)

    headers = ["#", "Portal", "Company", "Job Title", "Location",
               "Fit %", "Matched Skills", "Skill Gaps", "Verdict", "Apply Link"]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        cell.fill      = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = center
        cell.border    = border
    ws.row_dimensions[1].height = 28

    for i, r in enumerate(results, 2):
        fit      = r["fit"]
        row_fill = PatternFill("solid", fgColor=(
            "E2EFDA" if fit >= 85 else "FFEB9C" if fit >= 65 else "FFC7CE"
        ))
        row_data = [
            i - 1, r["portal"], r["company"], r["title"],
            r["location"], f"{fit}%", r["matches"], r["gaps"],
            r["verdict"], r["url"]
        ]
        for col, val in enumerate(row_data, 1):
            cell            = ws.cell(row=i, column=col, value=val)
            cell.fill       = row_fill
            cell.border     = border
            cell.font       = Font(name="Arial", size=9,
                                   color="0070C0" if col == 10 else "000000",
                                   underline="single" if col == 10 else None)
            cell.alignment  = center if col in (1, 6) else wrap

    for col, w in enumerate([4,10,22,28,16,7,35,30,40,50], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes  = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # Stats sheet
    ws2 = wb.create_sheet("Summary")
    stats = [
        ("Generated",             datetime.now().strftime("%d %b %Y, %H:%M")),
        ("Role",                  JOB_TITLE),
        ("Locations",             ", ".join(LOCATIONS)),
        ("Total Jobs Analyzed",   f"=COUNTA('Fitment Report'!B2:B1000)"),
        ("Strong Matches (≥85%)", f"=COUNTIF('Fitment Report'!F2:F1000,\">=85%\")"),
        ("Good Matches (65–84%)", f"=COUNTIFS('Fitment Report'!F2:F1000,\">=65%\",'Fitment Report'!F2:F1000,\"<85%\")"),
        ("Weak Matches (<65%)",   f"=COUNTIF('Fitment Report'!F2:F1000,\"<65%\")"),
    ]
    ws2["A1"]       = "AI Job Pipeline — Run Summary"
    ws2["A1"].font  = Font(bold=True, size=14, name="Arial", color="1F4E79")
    for row, (label, val) in enumerate(stats, 3):
        ws2.cell(row=row, column=1, value=label).font = Font(bold=True, name="Arial", size=10)
        ws2.cell(row=row, column=2, value=val).font   = Font(name="Arial", size=10)
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 32

    wb.save(path)
    return path

# ─── MAIN ────────────────────────────────────────────────────────────────────

def main():
    console.rule("[bold blue]AI Job Application Pipeline[/bold blue]")
    console.print(f"Role: [bold]{JOB_TITLE}[/bold]  |  Locations: [bold]{', '.join(LOCATIONS)}[/bold]\n")

    if ANTHROPIC_KEY == "YOUR_API_KEY_HERE":
        console.print("[red]ERROR:[/red] Set ANTHROPIC_KEY in your environment or in config.py")
        return

    client   = anthropic.Anthropic(api_key=ANTHROPIC_KEY)
    all_jobs = []
    per_loc  = max(1, MAX_JOBS // (2 * len(LOCATIONS)))

    # ── Step 1: Scrape ──────────────────────────────────────────────────────
    console.print("[bold cyan]Step 1 — Scraping job listings...[/bold cyan]")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page    = browser.new_context(user_agent=(
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/120.0.0.0 Safari/537.36"
        )).new_page()

        for loc in LOCATIONS:
            if SEARCH_LINKEDIN:
                all_jobs += scrape_linkedin(page, JOB_TITLE, loc, per_loc)
            if SEARCH_NAUKRI:
                all_jobs += scrape_naukri(page, JOB_TITLE, loc, per_loc)

        console.print(f"Found [bold]{len(all_jobs)}[/bold] listings — fetching JDs...")

        with Progress(SpinnerColumn(), TextColumn("{task.description}"), console=console) as prog:
            task = prog.add_task("", total=len(all_jobs))
            for job in all_jobs:
                prog.update(task, description=f"Fetching JD: {job['company'][:35]}")
                job["jd"] = fetch_jd(page, job["url"])
                prog.advance(task)
                time.sleep(0.4)

        browser.close()

    # ── Step 2: Fitment ─────────────────────────────────────────────────────
    console.print(f"\n[bold cyan]Step 2 — Claude fitment analysis ({len(all_jobs)} jobs)...[/bold cyan]")
    results = []

    with Progress(SpinnerColumn(), TextColumn("{task.description}"), console=console) as prog:
        task = prog.add_task("", total=len(all_jobs))
        for job in all_jobs:
            prog.update(task, description=f"Analyzing: {job['company'][:38]}")
            scores = run_fitment(client, job)
            if scores["fit"] >= MIN_FIT_PCT:
                results.append({**job, **scores})
            prog.advance(task)
            time.sleep(0.25)

    results.sort(key=lambda x: x["fit"], reverse=True)

    # ── Step 3: Print table ─────────────────────────────────────────────────
    console.print(f"\n[bold green]✓ {len(results)} jobs above {MIN_FIT_PCT}% threshold[/bold green]\n")
    tbl = Table(show_header=True, header_style="bold blue", box=None)
    for col, w in [("#",3),("Fit%",6),("Company",22),("Location",14),("Portal",9),("Verdict",48)]:
        tbl.add_column(col, width=w)
    for i, r in enumerate(results[:10], 1):
        fit_str = (
            f"[green]{r['fit']}%[/green]"   if r["fit"] >= 85 else
            f"[yellow]{r['fit']}%[/yellow]" if r["fit"] >= 65 else
            f"[red]{r['fit']}%[/red]"
        )
        tbl.add_row(str(i), fit_str, r["company"][:22], r["location"][:14],
                    r["portal"], r["verdict"][:48])
    console.print(tbl)

    # ── Step 4: Excel ───────────────────────────────────────────────────────
    out = f"job_fitment_report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    build_excel(results, out)
    console.print(f"\n[bold green]✓ Excel saved:[/bold green] {out}")

    # ── Step 5: Open top matches ────────────────────────────────────────────
    top = [r for r in results if r["fit"] >= 75][:TOP_N_OPEN]
    if top:
        console.print(f"\n[bold cyan]Opening top {len(top)} matches in browser...[/bold cyan]")
        for r in top:
            webbrowser.open(r["url"])
            time.sleep(0.8)

    console.rule("[bold blue]Done[/bold blue]")


if __name__ == "__main__":
    main()
